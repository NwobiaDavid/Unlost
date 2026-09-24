"""Turn a file on disk into LangChain Documents (one per page where pages exist), with OCR for scans and images."""

from __future__ import annotations

import logging
import os
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from langchain_core.documents import Document

log = logging.getLogger(__name__)

# OCR runs several Tesseract processes in parallel; one thread each avoids them fighting over cores.
os.environ.setdefault("OMP_THREAD_LIMIT", "1")

KINDS: dict[str, set[str]] = {
    "pdf": {".pdf"},
    "doc": {".docx"},
    "sheet": {".xlsx", ".xlsm", ".csv", ".tsv"},
    "slides": {".pptx"},
    "image": {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff", ".gif"},
    "text": {".txt", ".md", ".markdown", ".json", ".html", ".htm", ".xml", ".log", ".eml", ".rtf"},
}
EXT_TO_KIND = {ext: kind for kind, exts in KINDS.items() for ext in exts}
SUPPORTED = set(EXT_TO_KIND)

MAX_BYTES = 60 * 1024 * 1024
MAX_OCR_PAGES = 15
MAX_CHARS = 400_000  # per file; very long files are indexed up to this point


@dataclass
class Extracted:
    pages: list[Document] = field(default_factory=list)
    ocr: bool = False
    needs_ocr: bool = False  # OCR was skipped (fast first pass); run it later
    taken_at: datetime | None = None  # EXIF capture time for photos

    @property
    def text(self) -> str:
        return "\n".join(p.page_content for p in self.pages)


def kind_of(path: Path) -> str:
    return EXT_TO_KIND.get(path.suffix.lower(), "other")


def extract(path: Path, tesseract_cmd: str | None, ocr: bool = True) -> Extracted:
    """ocr=False skips the slow part (reading text from images and scans) and marks the result needs_ocr."""
    kind = kind_of(path)
    if kind == "pdf":
        return _pdf(path, tesseract_cmd, ocr)
    if kind == "image":
        return _image(path, tesseract_cmd, ocr)
    if kind == "doc":
        return _docx(path)
    if kind == "slides":
        return _pptx(path)
    if kind == "sheet":
        return _sheet(path)
    if kind == "text":
        return _plain(path)
    return Extracted()


def _pdf(path: Path, tesseract_cmd: str | None, ocr: bool = True) -> Extracted:
    try:
        pages = _pdf_text_pdfium(path)
    except Exception as e:  # pdfium is ~10x faster; fall back to pypdf for the odd file it can't open
        log.info("pdfium failed on %s (%s); trying pypdf", path, e)
        try:
            from pypdf import PdfReader

            pages = [Document(page_content=pg.extract_text() or "", metadata={"page": i})
                     for i, pg in enumerate(PdfReader(str(path)).pages, start=1)]
        except Exception as e2:  # encrypted or malformed PDFs
            log.info("pypdf failed on %s: %s", path, e2)
            pages = []
    text_chars = sum(len(d.page_content.strip()) for d in pages)
    # A scanned PDF has pages but (almost) no text layer: OCR it.
    scanned = not pages or text_chars < 40 * max(1, len(pages))
    if scanned and not ocr:
        return Extracted(pages=_cap(pages), needs_ocr=True)
    if tesseract_cmd and scanned:
        ocr_pages = _ocr_pdf(path, tesseract_cmd)
        if sum(len(d.page_content) for d in ocr_pages) > text_chars:
            return Extracted(pages=_cap(ocr_pages), ocr=True)
    return Extracted(pages=_cap(pages))


def _pdf_text_pdfium(path: Path) -> list[Document]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        pages, total = [], 0
        for i in range(len(pdf)):
            textpage = pdf[i].get_textpage()
            text = textpage.get_text_range()
            textpage.close()
            pages.append(Document(page_content=text, metadata={"page": i + 1}))
            total += len(text)
            if total > MAX_CHARS:
                break
        return pages
    finally:
        pdf.close()


def _ocr_pdf(path: Path, tesseract_cmd: str) -> list[Document]:
    import pypdfium2 as pdfium

    docs: list[Document] = []
    try:
        pdf = pdfium.PdfDocument(str(path))
    except Exception as e:
        log.info("pdfium failed on %s: %s", path, e)
        return docs
    try:
        for i in range(min(len(pdf), MAX_OCR_PAGES)):
            img = pdf[i].render(scale=2.0).to_pil()
            docs.append(Document(page_content=_ocr(img, tesseract_cmd), metadata={"page": i + 1}))
    finally:
        pdf.close()
    return docs


def _image(path: Path, tesseract_cmd: str | None, ocr: bool = True) -> Extracted:
    from PIL import Image, ImageOps

    with Image.open(path) as im:
        taken_at = _exif_date(im)
        if not ocr:
            return Extracted(taken_at=taken_at, needs_ocr=True)
        im = ImageOps.exif_transpose(im)
        im.thumbnail((2000, 2000))
        text = _ocr(im, tesseract_cmd) if tesseract_cmd else ""
    pages = [Document(page_content=text, metadata={"page": None})] if text.strip() else []
    return Extracted(pages=pages, ocr=bool(pages), taken_at=taken_at)


def _ocr(img, tesseract_cmd: str) -> str:
    import pytesseract

    pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    try:
        text = pytesseract.image_to_string(img, timeout=60)
    except (RuntimeError, pytesseract.TesseractError) as e:
        log.info("OCR failed: %s", e)
        return ""
    # Drop OCR noise: keep only lines containing a real word.
    lines = [ln.strip() for ln in text.splitlines() if re.search(r"[A-Za-z]{3,}|\d{2,}", ln)]
    return "\n".join(lines)


def _exif_date(im) -> datetime | None:
    try:
        exif = im.getexif()
        raw = exif.get_ifd(0x8769).get(36867) or exif.get(306)  # DateTimeOriginal, DateTime
        return datetime.strptime(str(raw).strip("\x00 "), "%Y:%m:%d %H:%M:%S") if raw else None
    except Exception:
        return None


def _docx(path: Path) -> Extracted:
    import docx2txt

    text = docx2txt.process(str(path)) or ""
    return Extracted(pages=_cap([Document(page_content=text, metadata={"page": None})]))


def _pptx(path: Path) -> Extracted:
    from pptx import Presentation

    pages = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        text = "\n".join(p for p in parts if p.strip())
        if text:
            pages.append(Document(page_content=text, metadata={"page": i}))
    return Extracted(pages=_cap(pages))


def _sheet(path: Path) -> Extracted:
    if path.suffix.lower() in (".csv", ".tsv"):
        return _plain(path)
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    pages, total = [], 0
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c).strip()]
                if cells:
                    rows.append(" | ".join(cells))
                    total += len(rows[-1])
                if total > MAX_CHARS:
                    break
            if rows:
                pages.append(Document(page_content=f"Sheet {ws.title}\n" + "\n".join(rows), metadata={"page": None}))
    finally:
        wb.close()
    return Extracted(pages=_cap(pages))


def _plain(path: Path) -> Extracted:
    raw = path.read_bytes()[: MAX_CHARS * 2]
    for enc in ("utf-8", "utf-16", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="ignore")
    suffix = path.suffix.lower()
    if suffix in (".html", ".htm", ".xml"):
        text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.S | re.I)
        text = re.sub(r"<[^>]+>", " ", text)
    elif suffix == ".rtf":
        text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    return Extracted(pages=[Document(page_content=text[:MAX_CHARS], metadata={"page": None})] if text.strip() else [])


def _cap(pages: list[Document]) -> list[Document]:
    out, total = [], 0
    for d in pages:
        if total >= MAX_CHARS:
            break
        d.page_content = d.page_content[: MAX_CHARS - total]
        total += len(d.page_content)
        if d.page_content.strip():
            out.append(d)
    return out


_MON = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
_MONTH_NUM = {m: i for i, m in enumerate(_MON, 1)}
_MON_RE = "|".join(_MON)
_DATE_PATTERNS = [
    # 2025-03-14, 2025/03/14
    (re.compile(r"\b(20\d\d|19\d\d)[-/.](0?[1-9]|1[0-2])[-/.](0?[1-9]|[12]\d|3[01])\b"),
     lambda m: (m[1], m[2])),
    # 14/03/2025 or 03/14/2025: day-first unless the middle number can't be a month
    (re.compile(r"\b(0?[1-9]|[12]\d|3[01])[-/.](0?[1-9]|[12]\d|3[01])[-/.](20\d\d|19\d\d)\b"),
     lambda m: (m[3], m[2] if int(m[2]) <= 12 else m[1])),
    # March 14, 2025 / 14 March 2025 / Mar 2025
    (re.compile(r"\b(?:\d{1,2}(?:st|nd|rd|th)?\s+)?(" + _MON_RE + r")[a-z]*\.?\s+(?:\d{1,2}(?:st|nd|rd|th)?,?\s+)?(20\d\d|19\d\d)\b", re.I),
     lambda m: (m[2], _MONTH_NUM[m[1].lower()[:3]])),
]


def months_mentioned(text: str, limit: int = 24) -> list[str]:
    """YYYY-MM values that appear in the text (invoice dates, statement periods...)."""
    found: dict[str, None] = {}
    for pattern, pick in _DATE_PATTERNS:
        for m in pattern.finditer(text[:100_000]):
            try:
                y, mo = (int(v) for v in pick(m))
            except (ValueError, KeyError):
                continue
            if 1 <= mo <= 12 and 1990 <= y <= 2100:
                found[f"{y:04d}-{mo:02d}"] = None
            if len(found) >= limit:
                return list(found)
    return list(found)
